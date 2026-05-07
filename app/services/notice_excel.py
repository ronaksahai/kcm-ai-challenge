"""
Notice Reply Drafting — Excel "Details Required" Generator
Produces a professional Excel sheet listing what documents/information
the client needs to provide, based on the notice analysis.

No external sample data dependency — all formatting is embedded here.
"""

import logging
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

# ── Styling Constants ────────────────────────────────────────

HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)

TITLE_FONT = Font(name="Calibri", size=14, bold=True, color="2F5496")
SUBTITLE_FONT = Font(name="Calibri", size=11, bold=False, color="4472C4")

DATA_FONT = Font(name="Calibri", size=11)
DATA_ALIGNMENT = Alignment(horizontal="left", vertical="top", wrap_text=True)
CENTER_ALIGNMENT = Alignment(horizontal="center", vertical="top", wrap_text=True)

THIN_BORDER = Border(
    left=Side(style="thin", color="B4C6E7"),
    right=Side(style="thin", color="B4C6E7"),
    top=Side(style="thin", color="B4C6E7"),
    bottom=Side(style="thin", color="B4C6E7"),
)

ALT_ROW_FILL = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")


# ── Main Generator ───────────────────────────────────────────

def generate_details_excel(notice_data: dict, output_path: str):
    """
    Generate the "Details Required" Excel workbook.

    Sheet structure:
    - Title block with assessee info
    - Table with columns: Sr. No | Point No. | Details / Documents Required | Remarks
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Details Required"

    # ── Title block ──────────────────────────────────────────
    assessee_name = notice_data.get("assessee_name", "____")
    pan = notice_data.get("pan", "____")
    ay = notice_data.get("assessment_year", "____")
    notice_type = notice_data.get("notice_type", "Notice")
    notice_date = notice_data.get("notice_date", "____")
    section = notice_data.get("notice_section")

    # Row 1: Title
    ws.merge_cells("A1:D1")
    title_cell = ws["A1"]
    title_cell.value = "Details Required from Client"
    title_cell.font = TITLE_FONT
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    # Row 2: Assessee info
    ws.merge_cells("A2:D2")
    info_cell = ws["A2"]
    section_ref = f" u/s {section}" if section else ""
    info_cell.value = (
        f"{assessee_name}  |  PAN: {pan}  |  A.Y. {ay}  |  "
        f"{notice_type}{section_ref} dated {notice_date}"
    )
    info_cell.font = SUBTITLE_FONT
    info_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 25

    # Row 3: Blank
    ws.row_dimensions[3].height = 10

    # ── Header row (Row 4) ───────────────────────────────────
    headers = ["Sr. No.", "Point No.", "Details / Documents Required", "Remarks"]
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGNMENT
        cell.border = THIN_BORDER

    ws.row_dimensions[4].height = 30

    # ── Data rows ────────────────────────────────────────────
    points = notice_data.get("points", [])

    for idx, point in enumerate(points):
        row_num = 5 + idx

        # Sr. No.
        cell_sr = ws.cell(row=row_num, column=1, value=idx + 1)
        cell_sr.font = DATA_FONT
        cell_sr.alignment = CENTER_ALIGNMENT
        cell_sr.border = THIN_BORDER

        # Point No.
        point_nums = point.get("point_numbers", str(idx + 1))
        cell_pt = ws.cell(row=row_num, column=2, value=point_nums)
        cell_pt.font = DATA_FONT
        cell_pt.alignment = CENTER_ALIGNMENT
        cell_pt.border = THIN_BORDER

        # Details / Documents Required
        details = point.get("details_for_client", "")
        if not details:
            # Fallback: use documents_asked list
            docs = point.get("documents_asked", [])
            if docs:
                details = "\n".join(f"• {d}" for d in docs)
            else:
                details = point.get("what_is_asked", "")
        cell_details = ws.cell(row=row_num, column=3, value=details)
        cell_details.font = DATA_FONT
        cell_details.alignment = DATA_ALIGNMENT
        cell_details.border = THIN_BORDER

        # Remarks
        remarks = point.get("ai_remarks", "")
        cell_remarks = ws.cell(row=row_num, column=4, value=remarks)
        cell_remarks.font = DATA_FONT
        cell_remarks.alignment = DATA_ALIGNMENT
        cell_remarks.border = THIN_BORDER

        # Alternating row color
        if idx % 2 == 1:
            for col in range(1, 5):
                ws.cell(row=row_num, column=col).fill = ALT_ROW_FILL

        # Auto row height (approximate)
        max_lines = max(
            len(details.split("\n")) if details else 1,
            len(remarks.split("\n")) if remarks else 1,
        )
        ws.row_dimensions[row_num].height = max(20, min(200, max_lines * 18))

    # ── Column widths ────────────────────────────────────────
    ws.column_dimensions["A"].width = 8    # Sr. No.
    ws.column_dimensions["B"].width = 12   # Point No.
    ws.column_dimensions["C"].width = 55   # Details Required
    ws.column_dimensions["D"].width = 40   # Remarks

    # ── Freeze panes ─────────────────────────────────────────
    ws.freeze_panes = "A5"

    # ── Print settings ───────────────────────────────────────
    ws.print_title_rows = "4:4"
    ws.sheet_properties.pageSetUpPr = None  # Reset

    # Save
    wb.save(output_path)
    logger.info(f"Details Required Excel saved to {output_path}")
